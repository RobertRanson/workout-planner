"""Convert the original Google Sheets workout tracker (.xlsx export) into data/seed.json.

Usage: python3 tools/import_xlsx.py Workout_Tracker.xlsx data/seed.json
"""
import datetime as dt
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# Sheet name -> (workout id, display name, calendar fill colour)
WORKOUTS = [
    ("Chest & Triceps", "chest", "Chest & Triceps", "FFFF0000"),
    ("Legs & Core", "legs1", "Legs & Core 1", "FF0000FF"),
    ("Back & Biceps", "back", "Back & Biceps", "FFFFFF00"),
    ("Legs & Core 2", "legs2", "Legs & Core 2", "FF00FFFF"),
    ("Upper Body", "upper", "Upper Body", "FF00FF00"),
]

MS = "https://www.muscleandstrength.com/exercises/"
# Links for exercises the sheet had no URL for (same exercise appears linked elsewhere, or standard M&S slug).
EXTRA_LINKS = {
    "Incline Dumbbell Bench Press": MS + "incline-dumbbell-bench-press.html",
    "Dumbbell Bench Press": MS + "dumbbell-bench-press.html",
}


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def parse_name(raw):
    """'Dumbbell Split Squat (8-12) Each' -> name, reps, each."""
    raw = re.sub(r"\s+", " ", str(raw)).strip()
    m = re.search(r"\(([^)]*)\)", raw)
    reps = None
    each = bool(re.search(r"\beach\b", raw, re.I))
    if m:
        inner = m.group(1)
        r = re.search(r"(\d+)\s*-\s*(\d+)", inner)
        if r:
            reps = [int(r.group(1)), int(r.group(2))]
        elif re.search(r"(\d+)\s*s\b", inner):
            reps = None
        else:
            n = re.search(r"\d+", inner)
            if n:
                reps = [int(n.group()), int(n.group())]
    name = re.sub(r"\([^)]*\)", "", raw)
    name = re.sub(r"\beach\b", "", name, flags=re.I)
    return re.sub(r"\s+", " ", name).strip(), reps, each


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, str) and re.fullmatch(r"\d{4}\.\d{2}\.\d{2}", v.strip()):
        return dt.datetime.strptime(v.strip(), "%Y.%m.%d").date()
    return None


def num(v):
    if v is None or v == "-" or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    return str(v).strip()


def find_row(ws, pred, max_row=6):
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            if pred(ws.cell(r, c).value):
                return r
    return None


def parse_sheet(ws):
    set_row = find_row(ws, lambda v: v == "Set 1")
    # Exercise names sit in the nearest row above "Set 1" holding text that isn't a URL.
    name_row = next(
        r for r in range(set_row - 1, 0, -1)
        if any(isinstance(v, str) and v.strip() and not v.startswith("http")
               for v in (ws.cell(r, c).value for c in range(2, ws.max_column + 1)))
    )
    unit_row = set_row + 1
    # Start column of each exercise = columns in name row with a non-link value.
    starts = []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(name_row, c).value
        if v and not str(v).startswith("http"):
            starts.append(c)
    exercises = []
    for i, c0 in enumerate(starts):
        c1 = (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_column
        raw = ws.cell(name_row, c0).value
        name, reps, each = parse_name(raw)
        link = None
        for r in range(1, set_row):
            cell = ws.cell(r, c0)
            if cell.hyperlink and cell.hyperlink.target:
                link = cell.hyperlink.target
            elif isinstance(cell.value, str) and cell.value.startswith("http"):
                link = cell.value
        link = link or EXTRA_LINKS.get(name)
        # Each set starts at a "Set N" header; its columns run until the next one.
        set_cols = [c for c in range(c0, c1 + 1) if str(ws.cell(set_row, c).value or "").startswith("Set")]
        sets = []
        for j, sc in enumerate(set_cols):
            end = (set_cols[j + 1] - 1) if j + 1 < len(set_cols) else c1
            cols = {}
            for c in range(sc, end + 1):
                u = str(ws.cell(unit_row, c).value or "").strip().lower()
                if u in ("weight", "reps", "time"):
                    cols[u] = c
            if not cols:  # e.g. Side Plank: "Time" sits one column right of "Set N"
                cols["time"] = sc + 1
            sets.append(cols)
        kind = "time" if all("time" in s and "weight" not in s for s in sets) else "weight"
        exercises.append({
            "id": slug(name),
            "name": name,
            "reps": reps,
            "each": each,
            "kind": kind,
            "sets": len(sets),
            "link": link,
            "_cols": sets,
        })
    sessions = []
    for r in range(unit_row + 1, ws.max_row + 1):
        d = parse_date(ws.cell(r, 1).value)
        if not d:
            continue
        entries = {}
        for ex in exercises:
            sets = []
            for cols in ex["_cols"]:
                if ex["kind"] == "time":
                    t = num(ws.cell(r, cols["time"]).value)
                    sets.append({"time": t} if t is not None else None)
                else:
                    w = num(ws.cell(r, cols["weight"]).value) if "weight" in cols else None
                    reps = num(ws.cell(r, cols["reps"]).value) if "reps" in cols else None
                    sets.append({"weight": w, "reps": reps} if (w is not None or reps is not None) else None)
            while sets and sets[-1] is None:
                sets.pop()
            if sets:
                entries[ex["id"]] = sets
        if entries:
            sessions.append({"date": d.isoformat(), "exercises": entries})
    for ex in exercises:
        del ex["_cols"]
    return exercises, sessions


def parse_calendar(ws):
    """Coloured cells in the Calendar sheet (week rows, Mon-Sun in E:K) -> [(date, workout id)]."""
    colour_to_id = {w[3]: w[1] for w in WORKOUTS}
    anchor_row, anchor = None, None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.startswith("=DATE("):
            y, m, d = map(int, re.findall(r"\d+", v)[:3])
            anchor_row, anchor = r, dt.date(y, m, d)
            break
    out = []
    header = find_row(ws, lambda v: v == "Mon", 20)
    for r in range(header + 1, ws.max_row + 1):
        week = anchor + dt.timedelta(weeks=r - anchor_row)
        v = ws.cell(r, 2).value
        if isinstance(v, dt.datetime):  # explicit date restarts the sequence
            break
        for i, c in enumerate(range(column_index_from_string("E"), column_index_from_string("K") + 1)):
            f = ws.cell(r, c).fill
            if f and f.fill_type and f.fgColor.type == "rgb" and f.fgColor.rgb in colour_to_id:
                out.append(((week + dt.timedelta(days=i)).isoformat(), colour_to_id[f.fgColor.rgb]))
    return out


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    workouts, sessions = [], []
    for sheet, wid, name, colour in WORKOUTS:
        exercises, ss = parse_sheet(wb[sheet])
        workouts.append({"id": wid, "name": name, "exercises": exercises})
        for s in ss:
            sessions.append({"id": f"{wid}-{s['date']}", "workout": wid, **s})
    # Calendar ticks without a matching logged session (within 3 days) become "done, no details" entries.
    calendar = parse_calendar(wb["Calendar"])
    print(f"{len(calendar)} calendar ticks")
    for date, wid in calendar:
        d = dt.date.fromisoformat(date)
        near = any(s["workout"] == wid and abs((dt.date.fromisoformat(s["date"]) - d).days) <= 3 for s in sessions)
        if not near:
            sessions.append({"id": f"{wid}-{date}", "workout": wid, "date": date, "exercises": {}})
    sessions.sort(key=lambda s: s["date"])
    with open(dst, "w") as fh:
        json.dump({"version": 1, "workouts": workouts, "sessions": sessions}, fh, indent=1)
    print(f"{len(workouts)} workouts, {len(sessions)} sessions -> {dst}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
