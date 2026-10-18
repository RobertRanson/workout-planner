"""Build data/videos.json from the video ID sheet(s).

Usage: python3 tools/import_videos.py data/videos.json Workout_Video_IDs.xlsx [Replacements.xlsx ...]

Every sheet with Workout (Day) / Exercise / Platform / Video ID columns is read; (Exercise) Source URL is
optional. Later files win, so a replacements sheet overrides earlier videos (rows without an ID are skipped).
Rows are matched to exercises in data/seed.json by Muscle & Strength URL, then by name, and the output
is keyed "workoutId/exerciseId" -> {"video": url, "link": M&S page}. A video also applies to the same
exercise (same M&S page) in other workouts.
"""
import json
import os
import re
import sys

import openpyxl

SEED = os.path.join(os.path.dirname(__file__), "..", "data", "seed.json")
def workout_id(day):
    d = re.sub(r"[^a-z0-9]+", " ", day.lower()).strip()
    if d.startswith("legs"):
        return "legs2" if d.endswith("2") else "legs1"
    for prefix in ("chest", "back", "upper"):
        if d.startswith(prefix):
            return prefix
    return None


def norm(name):
    name = re.sub(r"\([^)]*\)|\beach\b|\bdumbbells?\b", "", str(name), flags=re.I)
    return re.sub(r"[^a-z0-9]+", "", name.lower())


def rows(path):
    for ws in openpyxl.load_workbook(path).worksheets:
        header = [str(c.value or "").strip().lower() for c in ws[1]]
        idx = {h: i for i, h in enumerate(header)}
        src = idx.get("source url", idx.get("exercise source url"))
        day = idx.get("workout day", idx.get("workout"))
        if None in (day, idx.get("exercise"), idx.get("platform"), idx.get("video id")):
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            yield {
                "day": str(r[day] or ""),
                "name": str(r[idx["exercise"]] or ""),
                "platform": str(r[idx["platform"]] or "").strip().lower(),
                "id": str(r[idx["video id"]] or "").strip(),
                "link": str(r[src] or "").strip() if src is not None else "",
            }


def main(dst, *sources):
    workouts = json.load(open(SEED))["workouts"]
    out, unmatched = {}, []
    for i, src in enumerate(sources):
        replacing = i > 0
        for row in rows(src):
            apply(row, workouts, out, unmatched, replacing)
    finish(dst, workouts, out, unmatched)


def apply(row, workouts, out, unmatched, replacing):
    wid = workout_id(row["day"])
    if not wid or not row["id"]:
        return
    url = {"youtube": f"https://www.youtube.com/watch?v={row['id']}",
           "vimeo": f"https://vimeo.com/{row['id']}"}.get(row["platform"])
    if not url:
        return
    w = next(w for w in workouts if w["id"] == wid)
    ex = next((e for e in w["exercises"] if e["link"] and e["link"] == row["link"]), None) \
        or next((e for e in w["exercises"] if norm(e["name"]) == norm(row["name"])), None)
    if not ex:
        unmatched.append(f"{row['day']}: {row['name']}")
        return
    key = f"{wid}/{ex['id']}"
    if replacing:  # swap the video for this one exercise, keep its how-to link
        out[key] = {**out.get(key, {}), "video": url}
        return
    link = ex["link"] or (row["link"] if "muscleandstrength.com" in row["link"] else None)
    for w2 in workouts:
        for e2 in w2["exercises"]:
            if e2 is ex or (link and e2["link"] == link):
                out.setdefault(f"{w2['id']}/{e2['id']}", {"video": url, **({"link": link} if link else {})})
    out[key] = {"video": url, **({"link": link} if link else {})}


def finish(dst, workouts, out, unmatched):
    with open(dst, "w") as fh:
        json.dump(out, fh, indent=1, sort_keys=True)
    missing = [f"{w['id']}/{e['id']}" for w in workouts for e in w["exercises"] if f"{w['id']}/{e['id']}" not in out]
    print(f"{len(out)} exercises with video -> {dst}")
    print("unmatched rows:", unmatched or "none")
    print("exercises without video:", missing or "none")


if __name__ == "__main__":
    main(*sys.argv[1:])
